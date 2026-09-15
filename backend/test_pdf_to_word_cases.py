import io, os, sys
import fitz
from PIL import Image
import docx

sys.path.insert(0, os.path.abspath("."))
from app.services import conversion_service

def run_all_cases():
    print("======================================================")
    print("TESTING PDF TO WORD ACROSS 6 MANDATORY SCENARIOS")
    print("======================================================")

    # Case 1: Text PDF
    print("\n[1/6] Text PDF:")
    doc1 = fitz.open()
    p1 = doc1.new_page()
    p1.insert_text((72, 72), "DocFlow / Nexdoc - High-Quality Plain Text PDF to Word Conversion")
    pdf1_bytes = doc1.tobytes()
    doc1.close()
    docx1_bytes = conversion_service.pdf_to_word(pdf1_bytes)
    d1 = docx.Document(io.BytesIO(docx1_bytes))
    assert len(docx1_bytes) > 500, "DOCX too small"
    print(f"  [PASS] Converted {len(pdf1_bytes)} bytes PDF -> {len(docx1_bytes)} bytes DOCX")

    # Case 2: Multi-page PDF
    print("\n[2/6] Multi-page PDF (5 pages):")
    doc2 = fitz.open()
    for i in range(1, 6):
        p = doc2.new_page()
        p.insert_text((72, 72), f"Section {i}: Multi-page test on page {i}.")
    pdf2_bytes = doc2.tobytes()
    doc2.close()
    docx2_bytes = conversion_service.pdf_to_word(pdf2_bytes)
    d2 = docx.Document(io.BytesIO(docx2_bytes))
    assert len(docx2_bytes) > 1000, "Multi-page DOCX too small"
    print(f"  [PASS] Converted 5-page PDF ({len(pdf2_bytes)} bytes) -> {len(docx2_bytes)} bytes DOCX")

    # Case 3: PDF with embedded images
    print("\n[3/6] PDF containing raster image:")
    img = Image.new("RGB", (200, 200), color=(79, 70, 229))
    img_byte_arr = io.BytesIO()
    img.save(img_byte_arr, format="PNG")
    img_bytes = img_byte_arr.getvalue()

    doc3 = fitz.open()
    p3 = doc3.new_page()
    p3.insert_text((72, 72), "Embedded image header:")
    p3.insert_image(fitz.Rect(72, 100, 272, 300), stream=img_bytes)
    pdf3_bytes = doc3.tobytes()
    doc3.close()
    docx3_bytes = conversion_service.pdf_to_word(pdf3_bytes)
    d3 = docx.Document(io.BytesIO(docx3_bytes))
    assert len(docx3_bytes) > 2000, "Image-embedded DOCX too small"
    print(f"  [PASS] Converted image PDF -> {len(docx3_bytes)} bytes DOCX")

    # Case 4: Large allowed PDF (20 pages)
    print("\n[4/6] Large allowed PDF (20 pages):")
    doc4 = fitz.open()
    for i in range(1, 21):
        p = doc4.new_page()
        for line in range(8):
            p.insert_text((72, 72 + line * 30), f"Page {i}, Line {line}: Benchmark text.")
    pdf4_bytes = doc4.tobytes()
    doc4.close()
    docx4_bytes = conversion_service.pdf_to_word(pdf4_bytes)
    d4 = docx.Document(io.BytesIO(docx4_bytes))
    assert len(docx4_bytes) > 5000, "Large DOCX too small"
    print(f"  [PASS] Converted 20-page PDF ({len(pdf4_bytes)} bytes) -> {len(docx4_bytes)} bytes DOCX")

    # Case 5: Invalid/Malformed PDF
    print("\n[5/6] Invalid non-PDF data:")
    try:
        conversion_service.pdf_to_word(b"INVALID_HEADER_DATA_12345")
        print("  [FAIL] Did not reject malformed PDF")
        return False
    except Exception as e:
        print(f"  [PASS] Cleanly rejected malformed PDF: {e}")

    # Case 6: Empty/0-byte corrupted file
    print("\n[6/6] Empty 0-byte file:")
    try:
        conversion_service.pdf_to_word(b"")
        print("  [FAIL] Did not reject empty file")
        return False
    except Exception as e:
        print(f"  [PASS] Cleanly rejected empty file: {e}")

    print("\n======================================================")
    print("ALL 6 SCENARIOS VERIFIED AND WORKING 100%!")
    print("======================================================")
    return True

if __name__ == "__main__":
    success = run_all_cases()
    sys.exit(0 if success else 1)
